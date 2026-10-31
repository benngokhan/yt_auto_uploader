import contextlib
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import socket
import sys
import tempfile
import textwrap
import unittest
from unittest.mock import patch
import openpyxl

ROOT=Path(__file__).resolve().parent.parent
spec=importlib.util.spec_from_file_location('views_fix_draft',ROOT/'stats_views.py')
module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)

class ViewBaselineTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.work=Path(self.temp.name)
        self.table=self.work/'bilgidozu_yukleme_tablosu.xlsx'
        wb=openpyxl.Workbook();ws=wb.active;ws.title='DRIVE';ws.append(['h']*12)
        self.rows=[('video_1.mp4','aaa111aaa11',100),('video_2.mp4','bbb222bbb22',50),('video_3.mp4','ccc333ccc33',None),('video_4.mp4','ddd444ddd44',0),('video_5.mp4','eee555eee55',10),('video_6.mp4','fff666fff66','bad')]
        for name,vid,old in self.rows:
            ws.append(['',name,'Title '+name,'','','','','','','https://youtu.be/'+vid,'',old])
        wb.save(self.table);wb.close()
        self.values={'aaa111aaa11':120,'bbb222bbb22':45,'ccc333ccc33':514,'ddd444ddd44':5,'fff666fff66':7}
    def tearDown(self):self.temp.cleanup()
    def run_report(self,dry=False,fallback=None):
        prev=os.getcwd();os.chdir(self.work)
        try:
            def flat(tab):return self.values if tab=='shorts' else {}
            with patch.object(module,'TABLE',self.table),patch.object(module,'fetch_flat',side_effect=flat),patch.object(module,'fetch_one',side_effect=fallback or (lambda url: None)),patch.object(module,'run_ytdlp',side_effect=AssertionError('No subprocess')),patch.object(sys,'argv',['stats_views.py']+(['--dry-run'] if dry else [])):
                with open(os.devnull,'w') as sink, contextlib.redirect_stdout(sink):module.main()
            return json.loads((self.work/'stats_summary.json').read_text(encoding='utf8'))
        finally:os.chdir(prev)
    def test_blank_and_invalid_are_baselines_not_growth(self):
        s=self.run_report()
        self.assertEqual(s['total'],701)
        self.assertEqual(s['day_gain'],20)  # +20 -5 +5
        self.assertEqual(s['baseline_count'],2)
        self.assertEqual(s['compared_count'],3)
        self.assertEqual(s['unavailable_count'],1)
        self.assertEqual(s['updated'],5)
        self.assertFalse(s['dry_run'])
        self.assertEqual([(x['file'],x['delta']) for x in s['top']],[('video_1.mp4',20),('video_4.mp4',5)])
        wb=openpyxl.load_workbook(self.table,read_only=True,data_only=True);vals=[r[11] for r in list(wb.active.values)[1:]];wb.close()
        self.assertEqual(vals,[120,45,514,5,10,7])
    def test_dry_run_does_not_change_workbook(self):
        before=hashlib.sha256(self.table.read_bytes()).hexdigest();s=self.run_report(dry=True)
        self.assertEqual(hashlib.sha256(self.table.read_bytes()).hexdigest(),before)
        self.assertTrue(s['dry_run']);self.assertEqual(s['updated'],0);self.assertEqual(s['day_gain'],20)
    def test_all_missing_fails_closed_without_writing(self):
        self.values={};before=hashlib.sha256(self.table.read_bytes()).hexdigest()
        with self.assertRaisesRegex(RuntimeError,'Hiçbir video verisi'):
            self.run_report()
        self.assertEqual(hashlib.sha256(self.table.read_bytes()).hexdigest(),before)
        self.assertFalse((self.work/'stats_summary.json').exists())
    def test_no_linked_rows_fails_closed(self):
        wb=openpyxl.load_workbook(self.table)
        for row in range(2,8):wb.active.cell(row,10).value=None
        wb.save(self.table);wb.close()
        with self.assertRaisesRegex(RuntimeError,'Linkli video yok'):
            self.run_report()
        self.assertFalse((self.work/'stats_summary.json').exists())
    def test_fallback_attempt_counts_missing_not_row_index(self):
        wb=openpyxl.Workbook();ws=wb.active;ws.append(['h']*12)
        self.values={}
        for n in range(50):
            vid=f'id{n:09d}';ws.append(['',f'video_{n}.mp4','Title','','','','','','','https://youtu.be/'+vid,'',100])
            if n<49:self.values[vid]=100
        wb.save(self.table);wb.close();calls=[]
        def fallback(url):calls.append(url);return 105
        s=self.run_report(fallback=fallback)
        self.assertEqual(len(calls),1);self.assertIn('id000000049',calls[0])
        self.assertEqual(s['day_gain'],5);self.assertEqual(s['updated'],1);self.assertEqual(s['unavailable_count'],0)
    def test_telegram_message_is_escaped_and_dry_run_is_labeled(self):
        s={'total':50,'day_gain':-5,'updated':0,'baseline_count':1,'unavailable_count':2,'dry_run':True,'top':[{'delta':3,'title':'A & B <örnek>'}]}
        (self.work/'stats_summary.json').write_text(json.dumps(s,ensure_ascii=False),encoding='utf8')
        text=(ROOT/'.github/workflows/goruntulenme-raporu.yml').read_text(encoding='utf8')
        source=textwrap.dedent(text.split("          python - <<'PY'\n",1)[1].split('\n          PY',1)[0]);payload=[]
        class Reply:
            def read(self):return b'{}'
        def fake_urlopen(req,timeout=0):payload.append(json.loads(req.data.decode('utf8')));return Reply()
        prev=os.getcwd();os.chdir(self.work)
        try:
            with patch.dict(os.environ,{'TELEGRAM_BOT_TOKEN':'TEST_TOKEN','TELEGRAM_CHAT_ID':'TEST_CHAT','JOB_STATUS':'success','RUN_URL':'https://example.invalid'},clear=True),patch('urllib.request.urlopen',fake_urlopen):
                exec(compile(source,'workflow','exec'),{})
        finally:os.chdir(prev)
        self.assertEqual(len(payload),1);m=payload[0]['text']
        self.assertIn('-5',m);self.assertNotIn('+-5',m)
        self.assertIn('A &amp; B &lt;örnek&gt;',m)
        self.assertIn('İlk ölçüm (artışa dahil değil): 1',m)
        self.assertIn('Veri alınamayan: 2',m)
        self.assertIn('Test: tabloya yazılmadı',m)
        self.assertIn('Drive eşitlemesi bu raporla doğrulanmaz',m)

if __name__=='__main__':
    def deny(*a,**kw):raise AssertionError('network disabled')
    socket.create_connection=deny;socket.socket.connect=deny;socket.socket.connect_ex=deny
    unittest.main()
